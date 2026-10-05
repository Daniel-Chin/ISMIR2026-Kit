"""Load the private ConferenceCatalysts 2026 registration export.

This parser may need to be updated each year if the export format changes.
"""

import argparse
import datetime
from dataclasses import dataclass, field, fields
from pathlib import Path
from zoneinfo import ZoneInfo
from pprint import pprint

from dateutil.parser import parse as parse_datetime
from openpyxl import load_workbook


@dataclass
class Registrant:
    name: str = field(metadata={"header": "Full Name"})  # "Last, First"
    email: str = field(metadata={"header": "Email Address"})
    registration_type: str = field(metadata={"header": "Registration Type"})
    last_registration_date: datetime.datetime = field(
        metadata={
            "header": "Last Registration Date (GMT-05:00) Eastern [US & Canada]"
        }
    )
    virtual_tutorials: bool = field(metadata={"header": "Virtual Tutorials (Full Day)"})
    virtual_tutorials_students: bool = field(
        metadata={"header": "Virtual Tutorials (Full Day) - Students"}
    )
    tutorials: bool = field(metadata={"header": "Tutorials (Full Day)"})
    tutorials_students: bool = field(
        metadata={"header": "Tutorials (Full Day) - Students"}
    )

    def any_tutorials(self) -> bool:
        return self.virtual_tutorials or self.virtual_tutorials_students or self.tutorials or self.tutorials_students


@dataclass
class RawEntry(Registrant):
    invitee_status: str = field(metadata={"header": "Invitee Status"})

    def to_registrant(self) -> Registrant:
        assert self.invitee_status == "Accepted", "Invitee status must be Accepted"
        return Registrant(**{f.name: getattr(self, f.name) for f in fields(Registrant)})


def _parse_value(value: object, value_type: type) -> object:
    if value_type is bool:
        if value is None or value == "":
            return False
        if isinstance(value, str):
            value = value.strip().lower()
            if value in {"yes", "true", "1"}:
                return True
            if value in {"no", "false", "0", ""}:
                return False
        elif isinstance(value, (bool, int, float)) and value in (0, 1):
            return bool(value)
        raise ValueError("Expected a boolean, Yes/No, 1/0, or a blank cell")
    if value_type is datetime.datetime:
        if isinstance(value, str):
            value = parse_datetime(value)
        if not isinstance(value, datetime.datetime):
            raise ValueError("Expected a registration timestamp")
        if value.tzinfo is None:
            value = value.replace(tzinfo=ZoneInfo("America/New_York"))
        return value
    if not isinstance(value, str):
        raise TypeError("Expected a text cell")
    return value


def load_registration_data(file_path: str | Path) -> list[Registrant]:
    """Read row 4 headers and row 5 onward, validating each RawEntry.

    Blank tutorial selections mean false. Naive timestamps use US Eastern
    time, including daylight saving time. Unselected spreadsheet columns
    are ignored; wholly empty rows are skipped.
    """
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Registrant Details"]
        rows = sheet.iter_rows(min_row=4, values_only=True)
        headers = next(rows, ())
        columns = {}
        for f in fields(RawEntry):
            header = f.metadata["header"]
            if headers.count(header) != 1:
                raise ValueError(f"Expected exactly one column named {header!r}")
            columns[f.name] = headers.index(header)

        registrants = []
        for row_number, row in enumerate(rows, start=5):
            if all(value is None or value == "" for value in row):
                continue
            values = {}
            for f in fields(RawEntry):
                try:
                    values[f.name] = _parse_value(row[columns[f.name]], f.type)
                except (ValueError, TypeError, OverflowError):
                    # Do not include private cell contents in error messages.
                    raise ValueError(
                        f"Invalid {f.name} in spreadsheet row {row_number}"
                    ) from None
            entry = RawEntry(**values)
            registrants.append(entry.to_registrant())
        return registrants
    finally:
        workbook.close()


def load_registrants(argv: list[str] | None = None) -> list[Registrant]:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file_path", type=Path, help="Path to the private XLSX export")
    args = parser.parse_args(argv)
    return load_registration_data(args.file_path)


def main() -> None:
    input('Press Enter to view registrants with emails. Press Ctrl+C to abort.')
    registrants = load_registrants()
    for registrant in registrants:
        pprint(registrant)
        input('Enter...')


if __name__ == "__main__":
    main()
